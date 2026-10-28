#!/usr/bin/env python3
"""Build three focused Word guides from docs/user_guide/print.

Optional publishing tool; needs python-docx. It embeds the bundled OFL-licensed
application font. Render and inspect every output page before distribution.
This tool is separate from the standard-library skill runtime helper.
"""
import argparse
from datetime import datetime, timezone
from pathlib import Path
import re
import uuid
import zipfile
from lxml import etree
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
GUIDE = ROOT / 'docs/user_guide'
FONT = 'Google Sans Flex'
WIDTH = 6.614
GUIDES = [
    ('01-getting-started.md', 'BA-Mate-01-Getting-Started-v1.0.5.docx', 'Getting Started', 4),
    ('02-everyday-workflow.md', 'BA-Mate-02-Everyday-Workflow-v1.0.5.docx', 'Everyday Workflow', 5),
    ('03-testing-and-feedback.md', 'BA-Mate-03-Testing-and-Feedback-v1.0.5.docx', 'Testing and Feedback', 6),
]


def el(tag, attrs=None):
    x = OxmlElement(tag)
    for k, v in (attrs or {}).items():
        x.set(qn(k), str(v))
    return x


def set_font(style, size=11, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    for existing in list(rpr):
        if existing.tag in (qn('w:rFonts'), qn('w:color')):
            rpr.remove(existing)
    rpr.append(el('w:rFonts', {f'w:{k}': FONT for k in ('ascii', 'hAnsi', 'eastAsia', 'cs')}))
    rpr.append(el('w:color', {'w:val': '000000'}))


def link(p, label, target):
    h = el('w:hyperlink', {'r:id': p.part.relate_to(target, RT.HYPERLINK, is_external=True)})
    r = el('w:r'); props = el('w:rPr')
    props.append(el('w:rFonts', {'w:ascii': FONT, 'w:hAnsi': FONT}))
    props.append(el('w:color', {'w:val': '245E7F'}))
    props.append(el('w:u', {'w:val': 'single'}))
    r.append(props); t = el('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)


def inline(p, value):
    for chunk in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\))', value):
        if chunk.startswith('**'):
            p.add_run(chunk[2:-2]).bold = True
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)', chunk):
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', chunk)
            link(p, m[1], m[2])
        elif chunk:
            p.add_run(chunk)


def field(p, name, cached):
    f = el('w:fldSimple', {'w:instr': name})
    r = el('w:r'); t = el('w:t'); t.text = str(cached); r.append(t); f.append(r); p._p.append(f)


def picture(doc, path, alt, width=WIDTH):
    p = doc.add_paragraph(style='Figure')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('descr', alt)
    shape._inline.docPr.set('title', alt)


def table(doc, lines, reviewer=False):
    rows = [[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', c) for c in row)]
    count = len(rows[0]); t = doc.add_table(rows=1, cols=count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    if count == 3:
        if rows[0][0] == 'AI tool':
            ratios = [.18, .36, .46]
        else:
            ratios = [.35, .20, .45] if rows[0][0] in ('Your task', 'When you need to') else [.20, .39, .41]
    elif rows[0][0] in ('Dimension', 'Decision'):
        ratios = [.28, .72]
    else:
        ratios = [.27, .73]
    for col, ratio in zip(t.columns, ratios):
        col.width = Inches(WIDTH * ratio)
    props = t._tbl.tblPr
    borders = el('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(el('w:' + side, {'w:val': 'single', 'w:sz': 4, 'w:color': 'D9D9D9'}))
    props.append(borders)
    for i, row_values in enumerate(rows):
        row = t.rows[0] if i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr(); trpr.append(el('w:cantSplit'))
        if i == 0:
            trpr.append(el('w:tblHeader'))
        for j, value in enumerate(row_values):
            cell = row.cells[j]; cell.width = Inches(WIDTH * ratios[j]); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            mar = el('w:tcMar')
            for side, n in [('top', 110), ('bottom', 110), ('left', 135), ('right', 135)]:
                mar.append(el('w:' + side, {'w:w': n, 'w:type': 'dxa'}))
            tcpr.append(mar)
            tcpr.append(el('w:shd', {'w:fill': '193F5F' if i == 0 else ('F3F6F7' if i % 2 == 0 else 'FFFFFF')}))
            p = cell.paragraphs[0]; p.style = doc.styles['Table Text']; inline(p, 'Finding and evidence' if reviewer and i == 0 and j == 1 else ('' if reviewer and i > 0 and j == 1 else value))
            if reviewer and i > 0 and j == 1:
                p.style = doc.styles['Table Answer']
                response_control(p, 'review-' + str(i), rows[i][0] + ' finding', value)
            if count == 3 and j == 1 and rows[0][0] in ('Your task', 'When you need to'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                for r in p.runs:
                    r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255)
    # Add real paragraph spacing after a table without a visible empty line.
    p = doc.add_paragraph(style='Table Gap')
    return t


def setup(doc, number, short_title, pages):
    section = doc.sections[0]
    section.page_width = Mm(210); section.page_height = Mm(297)
    section.left_margin = Mm(21); section.right_margin = Mm(21)
    section.top_margin = Mm(18); section.bottom_margin = Mm(18)
    section.header_distance = Mm(8); section.footer_distance = Mm(8)
    section.different_first_page_header_footer = True
    for s in doc.styles:
        if s.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            set_font(s)
            ppr = s.element.find(qn('w:pPr'))
            if ppr is not None:
                for border in list(ppr.findall(qn('w:pBdr'))):
                    ppr.remove(border)
    for name, size, bold in [('Normal',11,False),('Title',27,True),('Subtitle',10.5,False),
                              ('Heading 1',17,True),('Heading 2',12,True),('Heading 3',11,True),
                              ('Header',8.5,False),('Footer',8.5,False)]:
        s=doc.styles[name]; set_font(s,size,bold)
        pf=s.paragraph_format; pf.line_spacing=1.14; pf.space_after=Pt(8); pf.widow_control=True
        if name.startswith('Heading'):
            pf.space_before=Pt(15); pf.space_after=Pt(7); pf.keep_with_next=True
        if name=='Title':
            pf.space_before=Pt(3); pf.space_after=Pt(7); pf.keep_with_next=True; pf.line_spacing=1.02
    for name,size,after in [('Prompt',11,10),('Table Text',10,0),('Figure',10,10),('Table Gap',1,8),('Guide Meta',9.5,11),('Prompt Label',9,3),('Form Label',11,3),('Form Hint',9,3),('Answer Line',11,2),('Table Answer',11,1),('Rating Choices',11,10)]:
        s=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH); s.base_style=doc.styles['Normal'];set_font(s,size)
        s.paragraph_format.space_after=Pt(after); s.paragraph_format.line_spacing=1.14
        if name=='Prompt':
            s.paragraph_format.left_indent=Mm(5);s.paragraph_format.right_indent=Mm(3)
            s.paragraph_format.keep_together=True
        if name=='Table Gap':
            s.paragraph_format.line_spacing=Pt(1)
    doc.styles['Figure'].paragraph_format.space_before=Pt(4)
    doc.styles['Figure'].paragraph_format.keep_together=True
    doc.styles['Prompt Label'].font.bold=True
    doc.styles['Prompt Label'].paragraph_format.left_indent=Mm(5)
    doc.styles['Prompt Label'].paragraph_format.space_before=Pt(3)
    doc.styles['Prompt Label'].paragraph_format.keep_with_next=True
    doc.styles['Form Label'].font.bold=True
    doc.styles['Form Label'].paragraph_format.space_before=Pt(8)
    doc.styles['Form Label'].paragraph_format.keep_with_next=True
    doc.styles['Form Hint'].font.color.rgb=RGBColor.from_string('595959')
    doc.styles['Form Hint'].paragraph_format.keep_with_next=True
    doc.styles['Answer Line'].paragraph_format.line_spacing=Pt(19)
    doc.styles['Table Answer'].paragraph_format.line_spacing=Pt(17)
    doc.styles['Rating Choices'].paragraph_format.space_before=Pt(3)
    doc.styles['Rating Choices'].paragraph_format.line_spacing=Pt(22)
    doc.styles['List Bullet'].paragraph_format.space_after=Pt(4)
    doc.styles['List Bullet'].paragraph_format.line_spacing=1.14
    if 'PlaceholderText' not in doc.styles:
        doc.styles.add_style('PlaceholderText',WD_STYLE_TYPE.CHARACTER)
    set_font(doc.styles['PlaceholderText'],10)
    doc.styles['PlaceholderText'].font.color.rgb=RGBColor.from_string('6C6C66')
    doc.styles['Subtitle'].paragraph_format.space_after=Pt(10)
    doc.styles['Footer'].paragraph_format.tab_stops.clear_all()
    header=section.header.paragraphs[0];header.style=doc.styles['Header']
    header.add_run(f'BA Mate   /   Guide {number}   /   {short_title}')
    for footer in (section.footer,section.first_page_footer):
        p=footer.paragraphs[0];p.style=doc.styles['Footer']
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH),WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('Research beta 1.0.5  ·  28 October 2026')
        p.add_run('\t');field(p,'PAGE',1);p.add_run(' / ');field(p,'NUMPAGES',pages)
    settings=doc.settings.element
    settings.append(el('w:updateFields',{'w:val':'true'}))
    settings.append(el('w:embedTrueTypeFonts'))
    settings.append(el('w:saveSubsetFonts',{'w:val':'false'}))
    cp=doc.core_properties;cp.author='BA Mate research project';cp.last_modified_by='BA Mate research project'
    cp.subject='Practical BA Mate research beta user guidance';cp.language='en-GB'
    cp.created=datetime(2026,10,26,tzinfo=timezone.utc);cp.modified=datetime(2026,10,28,tzinfo=timezone.utc)
    cp.keywords='BA Mate, business analysis, user guide, 1.0.5 beta'


def embed_fonts(path):
    """ECMA 376 font embedding; OFL and font name tables retain licence metadata."""
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    with zipfile.ZipFile(path) as z:
        data={name:z.read(name) for name in z.namelist()}
    ft=etree.fromstring(data['word/fontTable.xml'])
    for f in list(ft):
        if f.get(qn('w:name'))==FONT:ft.remove(f)
    f=el('w:font',{'w:name':FONT});f.append(el('w:family',{'w:val':'swiss'}));ft.append(f)
    relpath='word/_rels/fontTable.xml.rels'
    rels=etree.fromstring(data[relpath]) if relpath in data else etree.Element('{'+relns+'}Relationships',nsmap={None:relns})
    ids={r.get('Id') for r in rels}
    for style in ('Regular','Bold'):
        rid='rIdBAMate'+style
        if rid in ids:raise ValueError('Conflicting font relationship')
        key=uuid.uuid5(uuid.NAMESPACE_URL,'ba-mate/1.0.5/font/'+style)
        raw=bytearray((GUIDE/'fonts'/('GoogleSansFlex-'+style+'.ttf')).read_bytes())
        mask=key.bytes[::-1]
        for i in range(32):raw[i]^=mask[i%16]
        filename='GoogleSansFlex-'+style+'.odttf'
        data['word/fonts/'+filename]=bytes(raw)
        f.append(el('w:embed'+style,{'r:id':rid,'w:fontKey':'{'+str(key).upper()+'}'}))
        etree.SubElement(rels,'{'+relns+'}Relationship',Id=rid,Type=ns['r']+'/font',Target='fonts/'+filename)
    xml=lambda x:etree.tostring(x,xml_declaration=True,encoding='UTF-8',standalone=True)
    data['word/fontTable.xml']=xml(ft);data[relpath]=xml(rels)
    ct=etree.fromstring(data['[Content_Types].xml']);ctns='http://schemas.openxmlformats.org/package/2006/content-types'
    etree.SubElement(ct,'{'+ctns+'}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
    data['[Content_Types].xml']=xml(ct)
    tmp=path.with_suffix('.tmp')
    with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
        for name,value in data.items():z.writestr(name,value)
    tmp.replace(path)


CONTROL_ID = 1000


def control_properties(tag, label):
    global CONTROL_ID
    CONTROL_ID += 1
    props = el('w:sdtPr')
    props.append(el('w:id', {'w:val': CONTROL_ID}))
    props.append(el('w:tag', {'w:val': tag}))
    props.append(el('w:alias', {'w:val': label}))
    return props


def response_control(paragraph, tag, label, placeholder='Type here'):
    """Editable plain-text Word content control; lines also work in print."""
    sdt = el('w:sdt')
    props = control_properties(tag, label)
    props.append(el('w:text', {'w:multiLine': '1'}))
    props.append(el('w:showingPlcHdr'))
    sdt.append(props)
    content = el('w:sdtContent')
    r = el('w:r'); pr = el('w:rPr')
    pr.append(el('w:rStyle', {'w:val': 'PlaceholderText'})); r.append(pr)
    text = el('w:t'); text.text = placeholder; r.append(text); content.append(r)
    sdt.append(content); paragraph._p.append(sdt)


def answer_field(doc, tag, label, hint, lines):
    p = doc.add_paragraph(label, 'Form Label')
    if hint:
        doc.add_paragraph(hint, 'Form Hint')
    for i in range(int(lines)):
        p = doc.add_paragraph(style='Answer Line')
        if i == 0:
            response_control(p, tag, label)
        else:
            p.add_run(' ')
        p.paragraph_format.keep_with_next = i < int(lines) - 1
        ppr = p._p.get_or_add_pPr(); borders = el('w:pBdr')
        borders.append(el('w:bottom', {'w:val':'single','w:sz':4,'w:space':2,'w:color':'D9D9D9'}))
        borders.append(el('w:between', {'w:val':'single','w:sz':4,'w:space':2,'w:color':'D9D9D9'}))
        ppr.append(borders)


def checkbox(paragraph, tag, label):
    sdt = el('w:sdt'); props = control_properties(tag, label)
    box = el('w14:checkbox')
    box.append(el('w14:checked', {'w14:val':'0'}))
    box.append(el('w14:checkedState', {'w14:val':'2612','w14:font':'Segoe UI Symbol'}))
    box.append(el('w14:uncheckedState', {'w14:val':'2610','w14:font':'Segoe UI Symbol'}))
    props.append(box); sdt.append(props)
    content = el('w:sdtContent'); r = el('w:r'); pr = el('w:rPr')
    pr.append(el('w:rFonts',{'w:ascii':'Segoe UI Symbol','w:hAnsi':'Segoe UI Symbol'}))
    pr.append(el('w:sz',{'w:val':'26'})); r.append(pr)
    text = el('w:t'); text.text = '\u2610'; r.append(text)
    content.append(r); sdt.append(content); paragraph._p.append(sdt)


def choices(doc, tag, label, values, positions=None):
    p = doc.add_paragraph(label, 'Form Label')
    p.paragraph_format.space_before=Pt(10)
    p = doc.add_paragraph(style='Rating Choices')
    if positions is None:
        positions=[0,55,110,165,220,285,365]
    for n, (value, position) in enumerate(zip(values, positions)):
        if n:
            p.add_run('\t')
            p.paragraph_format.tab_stops.add_tab_stop(Pt(position))
        checkbox(p, tag + '-' + str(n+1), label + ': ' + value)
        p.add_run(' ' + value)
    p.paragraph_format.keep_together=True


def build(output):
    output.mkdir(parents=True,exist_ok=True)
    paths=[]
    for number,(source,name,title,pages) in enumerate(GUIDES,1):
        doc=Document();setup(doc,number,title,pages)
        picture(doc,GUIDE/'images/bamate_logo.png','BA Mate logo',width=1.0 if number>1 else 1.50)
        lines=(GUIDE/'print'/source).read_text().splitlines(); i=0; new_page=False; reviewer=False
        while i<len(lines):
            line=lines[i].strip();i+=1
            if not line:continue
            if line=='<!-- page -->':new_page=True;continue
            if line=='<!-- reviewer-table -->':reviewer=True;continue
            if line.startswith('<!-- field:'):
                args=line.removeprefix('<!-- field:').removesuffix('-->').strip().split('|')
                answer_field(doc,*args);continue
            if line.startswith('<!-- rating:'):
                tag,label,scale=line.removeprefix('<!-- rating:').removesuffix('-->').strip().split('|')
                choices(doc,tag,(tag+'  ' if tag.startswith('F') else '')+label,['1','2','3','4','5','NA','Skip']);continue
            if line.startswith('<!-- choice:'):
                parts=line.removeprefix('<!-- choice:').removesuffix('-->').strip().split('|')
                choices(doc,parts[0],parts[1],parts[2:],[0,108,193,275,369]);continue
            if line.startswith('|'):
                group=[line]
                while i<len(lines) and lines[i].strip().startswith('|'):
                    group.append(lines[i].strip());i+=1
                table(doc,group,reviewer);reviewer=False;continue
            if line.startswith('!['):
                m=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)
                picture(doc,(GUIDE/'print'/m[2]).resolve(),m[1]);continue
            if line.startswith('# '):p=doc.add_paragraph(line[2:],'Title');doc.core_properties.title=line[2:]
            elif line.startswith('## '):p=doc.add_paragraph(line[3:],'Heading 1')
            elif line.startswith('### '):p=doc.add_paragraph(line[4:],'Heading 2')
            elif line.startswith('Guide '):p=doc.add_paragraph(line,'Guide Meta')
            elif line.startswith('> '):
                doc.add_paragraph('PROMPT TO COPY', 'Prompt Label')
                p=doc.add_paragraph(style='Prompt');inline(p,line[2:])
            elif line.startswith('- '):
                p=doc.add_paragraph(style='List Bullet');inline(p,line[2:])
            else:
                prior = doc.paragraphs[-1] if doc.paragraphs else None
                p=doc.add_paragraph();inline(p,line)
                if prior is not None and prior.style.name == 'Answer Line':
                    p.paragraph_format.space_before=Pt(8)
            if new_page:
                p.paragraph_format.page_break_before=True
                p.paragraph_format.space_before=Pt(0)
                new_page=False
        path=output/name;doc.save(path);embed_fonts(path);paths.append(path)
    return paths


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,default=GUIDE)
    args=parser.parse_args()
    for path in build(args.output_dir):print(path)

if __name__=='__main__':main()
